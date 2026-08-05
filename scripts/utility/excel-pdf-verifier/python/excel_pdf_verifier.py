import pandas as pd
import pdfplumber
import re
from difflib import SequenceMatcher
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# =====================================================
# FILES
# =====================================================

EXCEL_FILE = r"SPA -NAME TENT CARD.xlsx"
PDF_FILE = r"NAme Tent card.pdf"

OUTPUT_FILE = "TentCard_Proofreading_Report.xlsx"

# =====================================================
# HELPERS
# =====================================================

def clean_text(text):
    if pd.isna(text):
        return ""

    text = str(text).upper()

    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^A-Z0-9&./ ]", "", text)

    return text.strip()


def similarity(a, b):
    return round(
        SequenceMatcher(
            None,
            clean_text(a),
            clean_text(b)
        ).ratio() * 100,
        2
    )


# =====================================================
# READ ALL SHEETS
# =====================================================

print("=" * 60)
print("READING EXCEL")
print("=" * 60)

xls = pd.ExcelFile(EXCEL_FILE)

all_data = []

for sheet in xls.sheet_names:

    print(f"Loading Sheet: {sheet}")

    df_sheet = pd.read_excel(
        EXCEL_FILE,
        sheet_name=sheet,
        dtype=str
    )

    df_sheet = df_sheet.fillna("")

    df_sheet["SOURCE_SHEET"] = sheet

    all_data.append(df_sheet)

df = pd.concat(
    all_data,
    ignore_index=True
)

original_columns = [
    c for c in df.columns
    if c != "SOURCE_SHEET"
]

print()
print("Sheets Found :", len(xls.sheet_names))
print("Rows Found   :", len(df))
print("Columns      :", len(original_columns))

# =====================================================
# EXTRACT PDF TEXT
# =====================================================

print()
print("=" * 60)
print("READING PDF")
print("=" * 60)

pdf_pages = []

with pdfplumber.open(PDF_FILE) as pdf:

    total_pages = len(pdf.pages)

    print("PDF Pages :", total_pages)

    for page_no, page in enumerate(pdf.pages, start=1):

        if page_no % 25 == 0:
            print(f"Processed {page_no}/{total_pages}")

        text = page.extract_text() or ""

        pdf_pages.append({
            "PDF_PAGE": page_no,
            "PDF_TEXT": clean_text(text)
        })

pdf_df = pd.DataFrame(pdf_pages)

# =====================================================
# MATCH EXCEL ROWS TO PDF PAGES
# =====================================================

print()
print("=" * 60)
print("VERIFYING DATA")
print("=" * 60)

results = []

total_rows = len(df)

for idx, row in df.iterrows():

    if (idx + 1) % 25 == 0:
        print(f"Checked {idx+1}/{total_rows}")

    row_text_parts = []

    for col in original_columns:

        value = clean_text(row[col])

        if value and value != "NAN":
            row_text_parts.append(value)

    search_text = " ".join(row_text_parts)

    best_page = None
    best_score = 0

    for _, pdf_row in pdf_df.iterrows():

        score = similarity(
            search_text,
            pdf_row["PDF_TEXT"]
        )

        if score > best_score:
            best_score = score
            best_page = pdf_row["PDF_PAGE"]

    if best_score >= 80:
        status = "MATCH"
    elif best_score >= 60:
        status = "CHECK"
    else:
        status = "MISMATCH"

    result = row.to_dict()

    result["BEST_PAGE"] = best_page
    result["MATCH_SCORE"] = best_score
    result["STATUS"] = status

    results.append(result)

# =====================================================
# CREATE REPORT
# =====================================================

print()
print("=" * 60)
print("CREATING REPORT")
print("=" * 60)

report_df = pd.DataFrame(results)

summary_df = (
    report_df
    .groupby(
        ["SOURCE_SHEET", "STATUS"]
    )
    .size()
    .reset_index(name="COUNT")
)

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    report_df.to_excel(
        writer,
        sheet_name="Proofreading",
        index=False
    )

    summary_df.to_excel(
        writer,
        sheet_name="Summary",
        index=False
    )

# =====================================================
# HIGHLIGHT RESULTS
# =====================================================

wb = load_workbook(OUTPUT_FILE)

ws = wb["Proofreading"]

yellow_fill = PatternFill(
    fill_type="solid",
    fgColor="FFF59D"
)

red_fill = PatternFill(
    fill_type="solid",
    fgColor="FF9999"
)

status_col = None

for cell in ws[1]:
    if cell.value == "STATUS":
        status_col = cell.column
        break

for row in range(2, ws.max_row + 1):

    status = ws.cell(
        row=row,
        column=status_col
    ).value

    if status == "CHECK":

        for col in range(1, ws.max_column + 1):
            ws.cell(row, col).fill = yellow_fill

    elif status == "MISMATCH":

        for col in range(1, ws.max_column + 1):
            ws.cell(row, col).fill = red_fill

wb.save(OUTPUT_FILE)

# =====================================================
# FINISHED
# =====================================================

match_count = (report_df["STATUS"] == "MATCH").sum()
check_count = (report_df["STATUS"] == "CHECK").sum()
mismatch_count = (report_df["STATUS"] == "MISMATCH").sum()

print()
print("=" * 60)
print("PROOFREADING COMPLETE")
print("=" * 60)

print("MATCH    :", match_count)
print("CHECK    :", check_count)
print("MISMATCH :", mismatch_count)

print()
print("OUTPUT FILE:")
print(OUTPUT_FILE)
